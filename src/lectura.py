"""Lectores sin limpieza: conservan valores y posiciones de las fuentes."""
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

FIELDS = {
    'ATM': ['id_transaccion', 'fecha_hora', 'id_cuenta', 'tipo_operacion', 'monto_mxn', 'id_cajero', 'resultado', 'comision_aplicada'],
    'DIGITAL': ['transfer_id', 'timestamp', 'account_id', 'amount', 'currency', 'beneficiary_clabe', 'notification_email', 'concept', 'status', 'platform'],
    'CATALOGO': ['id_cuenta', 'titular', 'tipo_cuenta', 'limite_credito_mxn', 'fecha_apertura', 'estado_cuenta'],
}


def leer_fuente(ruta, fuente):
    ruta = Path(ruta)
    if fuente not in FIELDS:
        raise ValueError('Fuente desconocida')
    raw = ruta.read_bytes()
    ext = ruta.suffix.lower()
    if fuente == 'DIGITAL' and ext == '.json':
        rows = json.loads(raw.decode('utf-8-sig'))
        if not isinstance(rows, list):
            raise ValueError('El JSON debe contener una lista')
    elif fuente in ('ATM', 'CATALOGO') and ext in ('.csv', '.xlsx'):
        if ext == '.csv':
            matrix = list(csv.reader(raw.decode('utf-8-sig').splitlines()))
        else:
            wb = load_workbook(ruta, read_only=True, data_only=False)
            try:
                if 'transacciones_atm' not in wb.sheetnames:
                    raise ValueError('Falta la hoja transacciones_atm')
                matrix = list(wb['transacciones_atm'].values)
            finally:
                wb.close()
        if not matrix:
            raise ValueError('Archivo sin encabezados')
        headers = [str(v).strip() if v is not None else '' for v in matrix[0]]
        if len(headers) != len(set(headers)) or not set(FIELDS[fuente]) <= set(headers):
            raise ValueError('Encabezados ausentes o duplicados')
        rows = [dict(zip(headers, row)) if len(row) == len(headers) else {'_fila_invalida': list(row)} for row in matrix[1:]]
    else:
        raise ValueError('Formato no admitido para esta fuente')
    return {'registros': rows, 'metadatos': {'archivo': ruta.name, 'hash': hashlib.sha256(raw).hexdigest(), 'fuente': fuente, 'offset': 1 if ext == '.json' else 2}}
